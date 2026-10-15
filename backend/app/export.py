import io
import random
from html import escape

import cairosvg
import docx
from weasyprint import HTML

from app.schemas import GradeWorksheet, QuestionType

_OPTION_LETTERS = "abcdefgh"


def _question_body_html(q, number: int) -> str:
    parts = [f'<p class="q"><b>{number}. ({q.marks} marks)</b> {escape(q.question)}</p>']

    if q.type == QuestionType.MCQ and q.options:
        options_html = "".join(
            f"<li>({_OPTION_LETTERS[i]}) {escape(opt)}</li>" for i, opt in enumerate(q.options)
        )
        parts.append(f'<ul class="options">{options_html}</ul>')

    elif q.type == QuestionType.TRUE_FALSE:
        parts.append('<p class="tf">[ &nbsp;&nbsp;] True &nbsp;&nbsp;&nbsp; [ &nbsp;&nbsp;] False</p>')

    elif q.type == QuestionType.MATCH_THE_FOLLOWING and q.match_pairs:
        lefts = [p.left for p in q.match_pairs]
        rights = [p.right for p in q.match_pairs]
        shuffled_rights = rights[:]
        random.shuffle(shuffled_rights)
        rows = "".join(
            f"<tr><td>{i + 1}. {escape(left)}</td><td>{chr(97 + i)}. {escape(right)}</td></tr>"
            for i, (left, right) in enumerate(zip(lefts, shuffled_rights))
        )
        parts.append(
            f'<table class="match"><tr><th>Column A</th><th>Column B</th></tr>{rows}</table>'
        )

    elif q.type == QuestionType.DIAGRAM_BASED and q.diagram_svg:
        parts.append(f'<div class="diagram">{q.diagram_svg}</div>')

    elif q.type == QuestionType.FILL_BLANK:
        pass  # the blank is already written into the question text itself

    elif q.type == QuestionType.SHORT_ANSWER:
        parts.append('<div class="answer-lines short"></div>')

    elif q.type == QuestionType.LONG_ANSWER:
        parts.append('<div class="answer-lines long"></div>')

    return "".join(parts)


_STYLE = """
body { font-family: 'Noto Sans', 'Lohit Devanagari', 'Lohit Tamil', sans-serif; padding: 30px; color: #1a1a1a; }
h1 { font-size: 20px; margin-bottom: 4px; text-align: center; }
.header-fields { display: flex; justify-content: space-between; font-size: 13px;
    border-top: 1px solid #333; border-bottom: 1px solid #333; padding: 6px 0; margin: 10px 0 20px; }
.meta { color: #333; font-size: 12px; margin-bottom: 20px; text-align: center; }
.page-break { page-break-before: always; }
.q { margin-bottom: 4px; }
.options { list-style: none; margin: 2px 0 16px 20px; padding: 0; }
.options li { margin-bottom: 3px; }
.tf { margin: 2px 0 16px 20px; }
.answer-lines { border-bottom: 1px solid #999; margin: 14px 0 16px 20px; }
.answer-lines.short { height: 34px; }
.answer-lines.long { height: 68px; }
.diagram { margin: 8px 0 16px 20px; }
.diagram svg { max-width: 260px; }
table.match { border-collapse: collapse; margin: 8px 0 16px 20px; width: 80%; }
table.match th, table.match td { border: 1px solid #999; padding: 5px 10px; text-align: left; font-size: 13px; }
table.match th { background: #f0f0f0; }
"""


def _worksheet_html(worksheet: GradeWorksheet, include_answers: bool) -> str:
    body = [_question_body_html(q, i) for i, q in enumerate(worksheet.questions, start=1)]

    answers_section = ""
    if include_answers:
        answer_rows = [f"<p>{i}. {escape(q.answer)}</p>" for i, q in enumerate(worksheet.questions, start=1)]
        answers_section = f"""
        <div class="page-break"></div>
        <h1>Answer Key — Grade {worksheet.grade}</h1>
        {"".join(answer_rows)}
        """

    header = "" if include_answers else f"""
    <div class="header-fields">
        <span>Name: ____________________</span>
        <span>Grade: {worksheet.grade}</span>
        <span>Date: ____________</span>
    </div>
    """

    return f"""
    <html>
    <head><style>{_STYLE}</style></head>
    <body>
        <h1>{escape(worksheet.title)}</h1>
        <p class="meta">Total marks: {worksheet.total_marks} &nbsp;|&nbsp; {escape(worksheet.instructions)}</p>
        {header}
        {"".join(body)}
        {answers_section}
    </body>
    </html>
    """


def worksheet_to_pdf(worksheet: GradeWorksheet, include_answers: bool = False) -> bytes:
    return HTML(string=_worksheet_html(worksheet, include_answers)).write_pdf()


def _add_docx_question(document: docx.Document, q, number: int) -> None:
    document.add_paragraph(f"{number}. ({q.marks} marks) {q.question}")

    if q.type == QuestionType.MCQ and q.options:
        for i, opt in enumerate(q.options):
            document.add_paragraph(f"    ({_OPTION_LETTERS[i]}) {opt}")

    elif q.type == QuestionType.TRUE_FALSE:
        document.add_paragraph("    [   ] True        [   ] False")

    elif q.type == QuestionType.MATCH_THE_FOLLOWING and q.match_pairs:
        lefts = [p.left for p in q.match_pairs]
        rights = [p.right for p in q.match_pairs]
        shuffled_rights = rights[:]
        random.shuffle(shuffled_rights)
        table = document.add_table(rows=len(lefts) + 1, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "Column A"
        table.cell(0, 1).text = "Column B"
        for row_i, (left, right) in enumerate(zip(lefts, shuffled_rights), start=1):
            table.cell(row_i, 0).text = f"{row_i}. {left}"
            table.cell(row_i, 1).text = f"{chr(97 + row_i - 1)}. {right}"

    elif q.type == QuestionType.DIAGRAM_BASED and q.diagram_svg:
        try:
            png_bytes = cairosvg.svg2png(bytestring=q.diagram_svg.encode(), output_width=350)
            document.add_picture(io.BytesIO(png_bytes))
        except Exception:
            if q.diagram_description:
                document.add_paragraph(f"    [Diagram: {q.diagram_description}]")

    elif q.type == QuestionType.SHORT_ANSWER:
        document.add_paragraph("    _______________________________________________")

    elif q.type == QuestionType.LONG_ANSWER:
        for _ in range(3):
            document.add_paragraph("    _______________________________________________")


def worksheet_to_docx(worksheet: GradeWorksheet, include_answers: bool = False) -> bytes:
    document = docx.Document()
    document.add_heading(worksheet.title, level=1)

    if not include_answers:
        header = document.add_table(rows=1, cols=3)
        header.cell(0, 0).text = "Name: ____________________"
        header.cell(0, 1).text = f"Grade: {worksheet.grade}"
        header.cell(0, 2).text = "Date: ____________"

    document.add_paragraph(f"Total marks: {worksheet.total_marks}\n{worksheet.instructions}")

    for i, q in enumerate(worksheet.questions, start=1):
        _add_docx_question(document, q, i)

    if include_answers:
        document.add_page_break()
        document.add_heading(f"Answer Key — Grade {worksheet.grade}", level=1)
        for i, q in enumerate(worksheet.questions, start=1):
            document.add_paragraph(f"{i}. {q.answer}")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
